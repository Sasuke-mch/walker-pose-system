"""Build the detailed DOCX from fixed wrist-processing experiment outputs."""
import argparse
from collections import Counter
from pathlib import Path
import shutil

import numpy as np

from experiment_wrist_processing import ROOT, read, sources, write, stats
from pose_app.fisheye_camera import load_stereo_fisheye
from pose_app.wrist_processing import perspective_view, upper_box, fisheye_rays, upright_to_raw

NAMES = {'full_flip': '原全身＋翻转', 'full_single': '原全身＋单次', 'upper_flip': '上肢裁剪＋翻转', 'upper_single': '上肢裁剪＋单次', 'perspective_flip': '局部透视＋翻转', 'perspective_single': '局部透视＋单次'}


def input_diagnosis(output):
    det, cached = sources()
    cal = load_stereo_fisheye(ROOT/'realtime_app/calibration/results')
    sparse = [__import__('json').loads(x) for x in (output/'sparse.jsonl').read_text().splitlines()]
    labels = read(ROOT/'research_records/engineering_validation/G20261003_wrist_detector_manual_v2/points.json')
    label_map = {(x['frame_index'], x['camera'], x['hand']): x for x in labels}
    records, coverage = [], {}
    for row in sparse:
        f, c, variant = row['frame'], row['camera'], row['variant']
        source = cached[c][f, :, :2]
        if variant.startswith('perspective') and not row['bootstrap_full']:
            prev = cached[c][f-1]
            K, D = (cal.K0, cal.D0) if c == 'left' else (cal.K1, cal.D1)
            view = perspective_view(upper_box(prev), K, D, c, support=prev[5:11, :2])
            rays = fisheye_rays(upright_to_raw(source, c), K, D) @ view.rotation
            q = rays @ view.intrinsic.T
            image_points = q[:, :2]/q[:, 2:3]
            front = rays[:, 2] > 0
        else:
            image_points, front = source, np.ones(17, bool)
        b = np.asarray(row['box'])
        center, scale = (b[:2]+b[2:])/2, (b[2:]-b[:2])*1.25
        scale = np.array([scale[0], scale[0]/.75]) if scale[0] > scale[1]*.75 else np.array([scale[1]*.75, scale[1]])
        input_xy = (image_points-center+scale/2)/scale*np.array([768, 1024])
        inside = front & (input_xy[:, 0] >= 0) & (input_xy[:, 0] < 768) & (input_xy[:, 1] >= 0) & (input_xy[:, 1] < 1024)
        coverage.setdefault(variant, []).extend(inside[11:17].astype(int).tolist())
        if variant == 'full_flip':
            for side, elbow, wrist in (('left', 7, 9), ('right', 8, 10)):
                label = label_map[f, c, side]
                if label['manual']['state'] != 'present':
                    continue
                p = np.array(row['raw'][wrist])
                e = np.array(row['raw'][elbow])
                direction = (p-e)/np.linalg.norm(p-e)
                delta = p-label['manual']['raw_px']
                records.append({'frame': f, 'camera': c, 'hand': side, 'error_px': float(np.linalg.norm(delta)), 'along_forearm_px': float(delta@direction), 'across_forearm_px': float(delta@np.array([-direction[1], direction[0]])), 'radial_pixel_distance': float(np.linalg.norm(p-((cal.K0 if c == 'left' else cal.K1)[:2, 2]))), 'input_forearm_px': row['input_forearm_length_px'][0 if side == 'left' else 1]})
    errors = [x['error_px'] for x in records]
    radial = [x['radial_pixel_distance'] for x in records]
    sizes = [x['input_forearm_px'] for x in records]
    result = {'labelled_points': records, 'lower_limb_input_coverage': {k: {'visible': sum(v), 'total': len(v), 'fraction': float(np.mean(v))} for k, v in coverage.items()}, 'radial_error_pearson': float(np.corrcoef(errors, radial)[0, 1]), 'input_forearm_error_pearson': float(np.corrcoef(errors, sizes)[0, 1]), 'along_forearm_abs_px': stats([abs(x['along_forearm_px']) for x in records]), 'across_forearm_abs_px': stats([abs(x['across_forearm_px']) for x in records]), 'along_forearm_negative_count': sum(x['along_forearm_px'] < 0 for x in records), 'measurement': 'forearm length and projection of detector residual, not observed wrist width, semantic bias proof or visibility truth'}
    write(output/'input_diagnosis.json', result)
    return result


def main():
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output', type=Path, required=True)
    ap.add_argument('--replace-report', action='store_true', help='replace this generated report and its backup')
    args = ap.parse_args()
    out = args.output.resolve()
    r = read(out/'results.json')
    depth = read(out/'da3_retry_v3/results.json')
    diagnosis = input_diagnosis(out)
    assets = out/'report_assets'
    assets.mkdir(exist_ok=True)
    path = out/'单模型腕部处理实验报告.docx'
    if path.exists() and not args.replace_report:
        raise FileExistsError(path)
    plt.rcParams['font.family'] = 'Microsoft YaHei'
    fig, ax = plt.subplots(1, 2, figsize=(11, 4))
    variants = ['full_flip', 'full_single', 'upper_single', 'perspective_single']
    ax[0].bar(range(4), [r['sparse'][v]['manual2d_px']['mean'] for v in variants], color='white', edgecolor='black')
    ax[0].set_xticks(range(4), ['原翻转', '原单次', '上肢单次', '透视单次'])
    ax[0].set_ylabel('人工腕点平均偏差 / px')
    ax[1].bar(range(4), [r['sparse'][v]['stereo_pair_processing_ms']['median'] for v in variants], color='white', edgecolor='black')
    ax[1].set_xticks(range(4), ['原翻转', '原单次', '上肢单次', '透视单次'])
    ax[1].set_ylabel('双目处理中位耗时 / ms')
    fig.tight_layout()
    fig.savefig(assets/'quality_time.png', dpi=180)
    plt.close(fig)
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5
    for h in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[h].font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
        doc.styles[h].element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    for s in doc.styles:
        for border in s.element.iter(qn('w:pBdr')):
            for edge in border:
                edge.set(qn('w:color'), '000000')
                for attr in ('themeColor', 'themeTint', 'themeShade'):
                    edge.attrib.pop(qn('w:'+attr), None)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.add_heading('单模型腕部处理实验报告', 0)
    doc.add_paragraph('2026 年 10 月 3 日　Sapiens2-0.4B；已有双目序列；不新增标注或采集')
    def para(t):
        doc.add_paragraph(t)
    def heading(t):
        doc.add_heading(t, 1)
    table_number = 0
    def table(headers, rows, caption):
        nonlocal table_number
        table_number += 1
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = value
        for values in rows:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value)
        for row in t.rows:
            for cell in row.cells:
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'FFFFFF')
                props.append(shade)
                borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    edge = OxmlElement('w:'+side)
                    for key, value in {'val': 'single', 'sz': '4', 'color': '000000'}.items():
                        edge.set(qn('w:'+key), value)
                    borders.append(edge)
                props.append(borders)
                for p in cell.paragraphs:
                    p.alignment = 1
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.size = Pt(9)
        para(f'表 {table_number}　{caption}')
    def fmt(value):
        return '—' if value is None else f'{value:.2f}'
    heading('一、摘要')
    para('完成单模型裁剪、透视、解码和后处理对照，并补充局部 DA3。连续透视的腕点均值从 25.67 降至 23.97 px，几何通过数下降。原生 FP32 双目仍约 1.2 s，暂不替换主线。')
    heading('二、数据与比较口径')
    para('数据来自同一段 448 帧双目序列。人工标注帧为 0、50、99、149、199、248、298、348、397、447；20 张图中 28 个腕点可评价，12 个槽位没有可用标注。8 组腕点具有双目人工参考，其中 5 组人工三角化通过原几何门。肘关节没有人工标注。')
    para('所有姿态候选只使用同一份 Sapiens2-0.4B 权重，输入 768×1024，不调用 PMPose、WiLoR 或新的人员检测器。全身对照沿用冻结检测框。局部输入依靠同一模型的肩、肘、腕，人工点只在评价阶段读取。历史全身扩框逻辑没有修改。')
    para('稀疏对照共 120 张配置图。连续对照取 10 个窗口，每个标注帧及其前 7 帧，第 0 帧窗口长度为 1；合计 73 对、每个候选 146 张图，三个单次版本共 438 次推理。每个窗口先全身初始化，随后完全使用自身上一帧预测，验证局部 ROI 的反馈影响。窗口较短，尚未覆盖长时间失跟与快速运动。')
    para('滤波和前臂约束重放完整 448 帧已保存的单模型观测。该组观测来自原翻转流程，因此后处理均与 26.33 px 原对照比较，不能与关闭翻转候选混用。')
    heading('三、输入、坐标与计时实现')
    para('原始鱼眼为 1920×1080，左右输入分别转正为 1080×1920。左视图回原始像素为 u=1919−y、v=x；右视图为 u=y、v=1079−x。Sapiens2 到 COCO17 的既有映射保持不变，腕来自原 308 点的左腕 62、右腕 41。所有候选在进入三角化前恢复为原始鱼眼像素。')
    para('上肢框用上一帧 COCO5..10 六个点，全部置信度至少 0.25；横向留 25% 余量，上方 30%、下方 45%，最小跨度 200×250 像素。框不可靠时回到全身检测，不用旧腕点填补当前结果。稀疏组读取已保存的同模型上一帧；连续组只能读取该候选自身上一帧。')
    para('透视展开使用真实 K、D，把上一帧上肢点转换成鱼眼射线，围绕这些射线建立虚拟相机，拟合局部范围并扩为 1.35 倍，保留 5% 余量。OpenCV 生成重采样表，模型在透视图上推理，再按虚拟射线投回真实鱼眼。虚拟相机不改变原标定、双目角色或三维单位。')
    para('首次试运行使用扩框背景网格确定透视范围，第 50 帧遇到跨虚拟相机后方的方向并中止，保留 v1 输出和日志。v2 改用六个上肢点作为几何支撑。两者不能合并为同一实验；v2 完整重跑所有配置。')
    para('本机为 RTX 5070 Ti Laptop 12 GB，PyTorch 2.11.0+cu128，原生 FP32、20 CPU 线程、TF32 关闭。计时前做 3 次预热，每次 CUDA 同步；计入视区生成、重采样、预处理、前向、GPU 到 CPU 传输、解码和逆映射。双目耗时为左右串行处理之和；不含采集、PNG 读取、人员检测、绘制及模型加载。未做 TensorRT、低精度或双目批处理，不能把这里的速度等同于它们的部署速度。')
    para(f'原翻转流程重新推理后的 28 个腕点与冻结结果逐点完全一致，最大差 {r["diagnosis"]["fresh_vs_cached_max_px"]:.8f} px。全部几何映射通过往返核对及独立 OpenCV map 测试。当前 2D 人工点和原结果没有重新标注。')
    heading('四、输入尺度诊断')
    table(['输入', '输入前臂均长 / px', '下肢点落入输入 / 120'], [[NAMES[v], fmt(r['sparse'][v]['input_forearm_px']['mean']), f'{diagnosis["lower_limb_input_coverage"][v]["visible"]}/120'] for v in ('full_single', 'upper_single', 'perspective_single')], '前臂尺度及完整身体信息代价')
    para('前臂尺度增加后，矩形裁剪并没有带来同等的腕点改善。说明当前误差不能全部归因于输入太小。下肢覆盖按同模型全身点是否落入局部输入计算，是输入可见范围检查；局部模型输出的有限腿点不等于腿部仍被图像观察。腕部实际宽度没有人工区域或分割依据，本次没有伪造该测量。')
    para(f'28 个标注点的腕误差与原始像素径向位置相关系数为 {diagnosis["radial_error_pearson"]:.3f}，与输入前臂长度为 {diagnosis["input_forearm_error_pearson"]:.3f}；沿前臂、横向绝对偏差均值为 {diagnosis["along_forearm_abs_px"]["mean"]:.2f}、{diagnosis["across_forearm_abs_px"]["mean"]:.2f} px。这些相关量来自同一段少量帧，不能据此确定畸变或腕点语义偏移就是原因。')
    para('28/28 个点的沿前臂有符号偏差均为负，即模型点相对于人工点偏向肘部。这比随机抖动更像持续的定位差异，应优先核对人工腕中心、模型腕定义和原图解码。当前没有发现左右翻转或像素恢复不一致，但尚不能确定是人工语义差异还是模型的域内偏移。不能直接按这 28 点拟合一个固定偏移再用同一批点宣称改善。')
    heading('五、裁剪、透视与翻转对照')
    table(['配置', '平均 / px', 'RMS / px', 'P95 / px', '双目中位 / ms'], [[NAMES[v], fmt(r['sparse'][v]['manual2d_px']['mean']), fmt(r['sparse'][v]['manual2d_px']['rms']), fmt(r['sparse'][v]['manual2d_px']['p95']), fmt(r['sparse'][v]['stereo_pair_processing_ms']['median'])] for v in ('full_flip', 'upper_flip', 'perspective_flip', 'full_single', 'upper_single', 'perspective_single')], '20 张标注图、28 个有效腕点；配置间使用同一组标注')
    doc.add_picture(str(assets/'quality_time.png'), width=Inches(6.1))
    para('图 1　腕点偏差与双目计算时间。图中速度不含采集和显示。')
    para('关闭翻转后，原全身误差从 26.33 降到 25.67 px，双目中位耗时从 2106.50 降到 1140.30 ms。这是当前最明确的计算量减少，但发生在这段数据的稀疏腕评价中，需要保留完整身体后续评价。矩形上肢裁剪没有稳定改善；透视单次为 23.74 px，比原单次下降约 7.5%。')
    table(['连续候选', '人工均值 / px', '诊断三维 / mm', '腕通过 / 146', '肘通过 / 146', '双目中位 / ms'], [[NAMES[v], fmt(x['manual2d_px']['mean']), fmt(x['stereo_diagnostic_mm']['mean']), f'{x["wrist_geometry_passed"]}/146', f'{x["elbow_geometry_passed"]}/146', fmt(x['stereo_pair_processing_ms']['median'])] for v, x in r['continuous'].items()], '73 对连续图；人工误差只用其中 28 个已标注点')
    para('连续透视候选的腕均值为 23.97 px，没有完全抵消稀疏收益；但腕、肘几何通过数都比全身单次低。全部 8 组有人工双目参考的腕点，三种候选均没有人工与模型同时通过的结果。诊断三维包含被拒绝点，不能将约 16 mm 写成经过验收的三维精度。透视候选目前不适合直接替换全身主线。')
    heading('六、UDP 解码与热图几何搜索')
    table(['配置', '原 UDP均值 / px', '整数峰值 / px', '热图几何搜索 / px'], [[NAMES[v], fmt(r['sparse'][v]['manual2d_px']['mean']), fmt(r['sparse'][v+'_argmax']['manual2d_px']['mean']), fmt(r['sparse'][v+'_epipolar']['manual2d_px']['mean'])] for v in ('full_flip', 'full_single', 'upper_single', 'perspective_single')], '同一份前向输出的解码和后处理比较')
    para('现有解码已使用 UDP/DARK，不需要再叠加一个亚像素算法。整数峰值在某些配置接近 UDP，但透视翻转中误差为 25.58 px，高于 UDP 的 24.10 px；没有依据删除原解码。')
    table(['单次输入', '原腕通过 / 20', '搜索后 / 20', '搜索后人工均值 / px', '人工与模型通过 / 8'], [[NAMES[v], r['sparse'][v]['wrist_geometry_passed'], r['sparse'][v+'_epipolar']['wrist_geometry_passed'], fmt(r['sparse'][v+'_epipolar']['manual2d_px']['mean']), r['sparse'][v+'_epipolar']['stereo_both_gates_mm']['n']] for v in ('full_single', 'perspective_single')], '局部热图约束提高几何一致性，却增加人工二维偏差')
    para('搜索限定原始像素移动 12 px，热图候选响应至少达到峰值的 0.5，只在峰值附近 5×5 网格选择；代价包括热图响应、鱼眼像素空间极线残差和移动量。原候选仍保留，不强行吸附极线。透视单次的腕通过数从 0/20 增加到 16/20，但人工误差从 23.74 增到 25.94 px。几何门改善和定位改善必须同时看，当前搜索不能作为默认定位修正。')
    heading('七、因果滤波与前臂约束')
    table(['滤波参数', '人工均值 / px', '加速度RMS / px·s⁻²', '平均位移 / px', '范围保留 / %', '每图中位 / ms'], [[m, fmt(x['manual2d_px']['mean']), fmt(x['acceleration_px_s2']['rms']), fmt(x['observation_shift_px']['mean']), fmt(x['coordinate_range_ratio']['mean']*100), f'{x["processing_ms"]["median"]:.4f}'] for m, x in r['filter'].items()], '完整 448 帧、双相机肘腕；原观测加速度 RMS 为 4899.71 px·s⁻²')
    para('滤波用真实序列时间戳，先低通速度，再按速度调节位置截止频率；缺失或不可靠观测直接重置并输出缺失，不用过去结果补点。三套固定配置均减少二阶变化，但腕人工均值均高于原 26.33 px，并减少部分坐标活动范围。当前宜作为可选显示稳定处理，原观测仍需保存。速度互相关最大值落在 0 帧，不代表滤波没有相位滞后或真实运动延迟。')
    f = r['forearm']
    counts = Counter(x['state'] for x in f['audit'])
    table(['项目', '数值'], [['通过肘腕联合几何的长度样本（左/右）', f'{f["accepted_length_samples"]["left"]}/{f["accepted_length_samples"]["right"]}'], ['实际更新手臂帧', f['updated_pairs']], ['人工腕均值 / px', fmt(f['manual2d_px']['mean'])], ['双目每帧中位耗时 / ms', fmt(f['processing_ms']['median'])], ['因几何/观测不更新', counts['not_updated_geometry_or_observation']]], '前臂长度约束的实际活动范围')
    para('前臂目标长度在至少 5 次可信观测后，用过去最多 30 次原始通过结果的中位数初始化。每帧联合优化肘、腕两点 XYZ，二维观测尺度 3 px、长度尺度 15 mm，最多 20 次函数评价；候选仍需正深度、12 px移动上限和原几何门。优化终止容差为 1e−5。只更新 34 个手臂帧，人工误差略增，不能用它替代腕部图像定位；其他失败状态均保留。')
    heading('八、人工参考敏感性')
    table(['假设坐标噪声σ / px', '平均空间变化 / mm', 'P95 / mm', '最大 / mm'], [[sigma, fmt(x['mean']), fmt(x['p95']), fmt(x['max'])] for sigma, x in r['diagnosis']['manual_sensitivity_mm'].items()], '8 组参考，每组每尺度 30 次独立高斯扰动，seed0')
    para('2 px 假设噪声对应平均 2.27 mm、P95 4.64 mm 的空间变化；5 px 对应平均 5.59 mm、P95 12.22 mm。本次透视输入的三维均值改善约 1 mm，尚不能据此认定空间精度已经改善。这个试验测的是条件敏感性，实际人工重复标注误差尚未测量。')
    heading('九、局部 DA3 诊断')
    table(['项目', '结果'], [['双目推理中位 / ms', fmt(depth['paired_inference_ms']['median'])], ['双目含透视处理中位 / ms', fmt(depth['paired_with_view_ms']['median'])], ['单目推理中位 / ms', fmt(depth['monocular_inference_ms']['median'])], ['深度有限比例', '100%'], ['肘腕采样可用', f'{depth["available_samples"]}/{depth["total_samples"]}'], ['被深度修改的姿态点', depth['pose_points_modified']]], 'DA3-LARGE-1.1：10 次双目及 20 次单目，process_res504')
    para('DA3 用同源局部透视图，提供虚拟内参，不提供外参或米制对齐。直接复用官方输入处理、模型前向和输出转换，保留配对与单目两种结果。深度及置信度采样使用官方预处理后的内参定位，不把原图像素直接索引缩小的深度图。输出为 504×378 局部相对深度，160 个肘腕采样均可用。')
    para('仓库 AGENTS.md 规定 DA3 只用于图像空间 ROI/遮挡诊断，禁止用于标定、人物关联、关键点筛选、三角化和米制深度。本轮因此没有执行 DA3 深度作为三维损失或米制腕坐标的实验，不能把有限深度图写成腕点精修通过。它提供的是局部层次信息，还缺少遮挡真值与腕部几何监督。')
    para('旧 DA3 虚拟环境的 Python 基路径失效。缺少的依赖装入 outputs/da3_dependencies_20261003，主环境未更换。首次调用漏传官方前向需要的空辅助层列表，第二次发现旧版输出转换器会将缺失 is_metric 读取成空 Dict；两次失败目录保留。第三次按原输出字典取标记，无标记记为 null，30 次正式调用全部完成。')
    para('局部 DA3 额外增加约 218 ms 双目处理，即使以后姿态网络加速，这项成本也仍需单独评估。当前更适合作为离线局部诊断或低频 ROI 候选研究。反畸变来自已有标定，完全可以先独立完成，不依赖 DA3。')
    doc.add_picture(str(out/'da3_retry_v3/0199_1_depth.png'), width=Inches(6.1))
    para('图 2　第 199 帧右相机局部透视图与 DA3 相对深度。颜色显示该图内部相对层次，不标米制距离。')
    heading('十、决定与尚未完成的验证')
    table(['处理', '当前决定', '原因'], [['关闭翻转', '优先进入部署候选', '本段腕误差不增且计算约减半；完整身体仍须复核'], ['普通上肢裁剪', '不设默认', '分辨率提高但腕误差未稳定改善'], ['局部透视', '保留候选', '腕误差降低，肘腕几何通过变差且下肢覆盖减少'], ['UDP/DARK', '保持原解码', '已实现；整数峰值无稳定优势'], ['热图几何搜索', '不设默认', '通过率增加，但人工定位偏差变大'], ['因果滤波', '可选显示/噪声处理', '低成本，存在定位偏差和动作幅度取舍'], ['前臂长度约束', '诊断候选', '活动样本少，腕误差没有改善'], ['局部DA3', '离线或低频诊断', '额外模型成本及相对深度证据不足']], '本轮测试后的方法决定，不升级主线接受级别')
    para('实际实时瓶颈是原生 FP32 Sapiens2 前向。双目单次处理仍约 1.14～1.18 s，30 FPS 对应每对总预算 33.3 ms，现有配置差距很大。本轮没有测低精度、TensorRT、批处理或更小输入，不能给它们编造 FPS。下一步应先固定全身单次候选，单独比较推理精度与部署优化，再决定局部透视能否以低频补充进入完整身体系统。')
    para('本轮没有肘部人工标签，肘几何通过数只能检查跨视角解释是否一致；没有完成长序列局部闭环失跟测试或真实设备采集计时。局部视图虽然每帧只用一个姿态模型，也会丢失部分全身观测，不能直接作为完整 SMPL-H 主线输入。')
    heading('十一、代码、复现与验证')
    for command in [
        'python -B realtime_app/experiment_wrist_processing.py prepare --output <新目录>',
        'python -B realtime_app/experiment_wrist_processing.py run --output <同目录>',
        'python -B realtime_app/experiment_wrist_processing.py analyze --output <同目录>',
        'python -B realtime_app/experiment_wrist_depth.py --output <同目录> --run-name da3_retry_v3 --da3-root <已有DA3源码> --model-dir <DA3-LARGE-1.1> --dependencies outputs/da3_dependencies_20261003',
        'python -B realtime_app/report_wrist_processing.py --output <同目录>']:
        para(command)
    para('需使用已有 .venv-cuda 环境。报告脚本读取 da3_retry_v3/results.json。输入路径与参数见 protocol.json、postprocessing_protocol.json 及 DA3 协议；逐帧预测见 sparse.jsonl、continuous.jsonl；原始拒绝及统计见 results.json；输入诊断见 input_diagnosis.json。报告及相关图保存在当前实验目录，实验报告文件夹保存同名备份。')
    para('新增测试覆盖左右旋转、OpenCV 重采样映射、虚拟后方拒绝、低置信度 ROI、滤波因果性、缺失不补点、图像支持候选和前臂更新移动限制。原全身冻结腕点数值复现完全一致。公开测试、linter、无缓存编译及文档结构/渲染检查结果另存 verification.json；现有拟合求解器和方法级别未改。')
    if (out/'verification.json').exists():
        v = read(out/'verification.json')['public_tests']
        para(f'最终公开验证：{v["passed"]} 项通过、{v["skipped"]} 项明确跳过、{v["subtests_passed"]} 个子测试通过；sequence_pipeline 另 {v["sequence_passed"]} 项通过。新增局部模块共 8 项正反例测试。')
    heading('十二、参考资料')
    for t in ['Sapiens2官方代码与UDP热图实现：third_party/sapiens2/repo/sapiens/pose/src/datasets/codecs/udp_heatmap.py', 'ZoomNAS：https://arxiv.org/abs/2208.11547', 'Epipolar Transformers：https://arxiv.org/abs/2005.04551', 'Ego-Exo 3D Hand Tracking：https://hands-workshop.org/files/2025/Egocentric.pdf', 'Depth Anything 3官方源码：https://github.com/ByteDance-Seed/Depth-Anything-3']:
        para(t)
    doc.save(path)
    backup = ROOT/'资料/实验报告/单模型腕部处理实验报告_20261003.docx'
    if backup.exists() and not args.replace_report:
        raise FileExistsError(backup)
    shutil.copyfile(path, backup)
    print(path)


if __name__ == '__main__':
    main()
