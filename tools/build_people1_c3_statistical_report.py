from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "research_records" / "reports" / "20260906_people1_c3_statistical_report"
ASSET_DIR = REPORT_DIR / "report_assets"
DOCX_PATH = REPORT_DIR / "people1_c3_人体数据简要统计报告.docx"

FONT_PATH = Path(r"C:\Windows\Fonts\simhei.ttf")

RUNS = {
    "people_1": {
        "label": "people_1（远到近）",
        "root": ROOT / "research_records" / "engineering_validation" / "V20260906_people1_30fps_c3_2d_3d",
    },
    "people_1_near": {
        "label": "people_1_near（近距离）",
        "root": ROOT / "research_records" / "engineering_validation" / "V20260906_people1_near_30fps_c3_fullchain",
    },
}


def font(size: int):
    return ImageFont.truetype(str(FONT_PATH), size=size)


def text_size(draw: ImageDraw.ImageDraw, value: str, fnt):
    box = draw.textbbox((0, 0), value, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def load_2d(root: Path, model: str):
    with (root / f"{model}_vs_sapiens2_2d_coco17" / "summary_overall.csv").open(encoding="utf-8-sig", newline="") as handle:
        overall = next(csv.DictReader(handle))
    with (root / f"{model}_vs_sapiens2_2d_coco17" / "summary_by_body_region.csv").open(encoding="utf-8-sig", newline="") as handle:
        regions = {row["body_region"]: row for row in csv.DictReader(handle)}
    return {
        "mean": float(overall["mean_relative_error_px"]),
        "pck25": 100.0 * float(overall["relative_pck_25px"]),
        "regions": {key: float(regions[key]["mean_relative_error_px"]) for key in ("head", "shoulder_girdle", "upper_limb", "lower_limb")},
    }


def load_3d(root: Path, model: str):
    summary = load_json(root / f"{model}_ungated_stereo" / "offline_stereo_summary.json")
    t1 = load_json(root / f"{model}_ungated_lower_limb_t1_t4" / "t1_trajectory" / "trajectory_summary.json")
    t2 = load_json(root / f"{model}_ungated_lower_limb_t1_t4" / "t2_kinematics" / "kinematics_summary.json")
    return {
        "pairs": int(summary["processed_pairs"]),
        "matched": int(summary["matched_pairs"]),
        "points": int(summary["total_valid_3d_keypoints"]),
        "lower_coverage": 100.0 * float(t1["observed_lower_limb_coverage"]),
        "missing": summary.get("matched_keypoint_rejection_reasons", {}),
        "flags": summary.get("triangulated_quality_flags", {}),
        "t2": {name: value for name, value in t2["per_metric"].items()},
    }


def plot_2d(rows: list[dict], output: Path):
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    title_font, body_font, small_font = font(36), font(25), font(20)
    draw.text((80, 35), "图 1  C3 自适应连续扩框的二维结果", fill="black", font=title_font)
    colors = {"pmpose": "#1769AA", "probpose": "#E67E22"}
    panels = [
        ("相对 Sapiens2 平均像素差 (px)", "mean", 32.0, 160),
        ("PCK@25 (%)", "pck25", 100.0, 570),
    ]
    for panel_title, key, max_value, top in panels:
        left, right, bottom = 130, 1710, top + 280
        draw.text((left, top - 45), panel_title, fill="black", font=body_font)
        for step in range(5):
            y = bottom - step * (bottom - top) / 4
            value = max_value * step / 4
            draw.line((left, y, right, y), fill="#D9D9D9", width=2)
            label = f"{value:.0f}"
            draw.text((55, y - 12), label, fill="#555555", font=small_font)
        group_x = [480, 1240]
        for group, video_key in enumerate(("people_1", "people_1_near")):
            values = [row for row in rows if row["video_key"] == video_key]
            for offset, row in zip((-85, 85), values):
                height = (bottom - top) * row[key] / max_value
                x0, x1 = group_x[group] + offset - 55, group_x[group] + offset + 55
                y0 = bottom - height
                draw.rectangle((x0, y0, x1, bottom), fill=colors[row["model"]])
                value_label = f"{row[key]:.2f}" if key == "mean" else f"{row[key]:.1f}"
                width, _ = text_size(draw, value_label, small_font)
                draw.text((group_x[group] + offset - width / 2, y0 - 28), value_label, fill="black", font=small_font)
            label = "people_1\n远到近" if video_key == "people_1" else "people_1_near\n近距离"
            draw.multiline_text((group_x[group] - 70, bottom + 18), label, fill="black", font=small_font, spacing=4, align="center")
    draw.rectangle((1390, 52, 1420, 78), fill=colors["pmpose"])
    draw.text((1430, 52), "PMPose", fill="black", font=small_font)
    draw.rectangle((1580, 52, 1610, 78), fill=colors["probpose"])
    draw.text((1620, 52), "ProbPose", fill="black", font=small_font)
    image.save(output)


def read_kinematic_series(root: Path, model: str):
    path = root / f"{model}_ungated_lower_limb_t1_t4" / "t2_kinematics" / "lower_limb_kinematics.csv"
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            records.append({
                "time": float(row["pair_timestamp_sec"]),
                "left": float(row["left_knee_angle_deg"]) if row["left_knee_angle_deg_available"] == "True" else None,
                "right": float(row["right_knee_angle_deg"]) if row["right_knee_angle_deg_available"] == "True" else None,
            })
    return records


def draw_curve_panel(draw, rect, records, title):
    left, top, right, bottom = rect
    body_font, small_font = font(25), font(19)
    draw.text((left, top - 38), title, fill="black", font=body_font)
    ymin, ymax = 100.0, 185.0
    for step in range(5):
        y = bottom - step * (bottom - top) / 4
        value = ymin + (ymax - ymin) * step / 4
        draw.line((left, y, right, y), fill="#E0E0E0", width=2)
        draw.text((left - 68, y - 10), f"{value:.0f}", fill="#555555", font=small_font)
    max_time = max(item["time"] for item in records)
    for step in range(5):
        x = left + step * (right - left) / 4
        value = max_time * step / 4
        draw.line((x, top, x, bottom), fill="#F0F0F0", width=1)
        draw.text((x - 12, bottom + 10), f"{value:.1f}", fill="#555555", font=small_font)
    draw.text((right - 90, bottom + 45), "时间 (s)", fill="#555555", font=small_font)
    for key, color in (("left", "#1769AA"), ("right", "#D81B60")):
        previous = None
        for item in records:
            value = item[key]
            if value is None:
                previous = None
                continue
            x = left + item["time"] / max_time * (right - left)
            y = bottom - (value - ymin) / (ymax - ymin) * (bottom - top)
            if previous is not None:
                draw.line((previous[0], previous[1], x, y), fill=color, width=2)
            previous = (x, y)
    draw.line((left, top, left, bottom), fill="black", width=2)
    draw.line((left, bottom, right, bottom), fill="black", width=2)


def plot_knee_curves(series: dict[str, list[dict]], output: Path):
    image = Image.new("RGB", (1800, 1150), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 35), "图 2  PMPose 无门限三角化后的左右膝角时序", fill="black", font=font(36))
    draw_curve_panel(draw, (150, 145, 1710, 515), series["people_1"], "people_1（远到近）")
    draw_curve_panel(draw, (150, 690, 1710, 1060), series["people_1_near"], "people_1_near（近距离）")
    draw.rectangle((1280, 70, 1310, 94), fill="#1769AA")
    draw.text((1320, 68), "左膝", fill="black", font=font(20))
    draw.rectangle((1440, 70, 1470, 94), fill="#D81B60")
    draw.text((1480, 68), "右膝", fill="black", font=font(20))
    image.save(output)


def set_cell_shading(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_run_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold


def add_paragraph(doc, text="", *, size=10.5, bold=False, align=None, before=0, after=4):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.35
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "FFFFFF")
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(header)
        set_run_font(run, size=9, bold=True)
        run.font.color.rgb = RGBColor(0, 0, 0)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if widths:
                cells[index].width = Cm(widths[index])
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cells[index], "FFFFFF")
            para = cells[index].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(str(value))
            set_run_font(run, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def main():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    rows = []
    body = {}
    for video_key, config in RUNS.items():
        body[video_key] = {}
        for model in ("pmpose", "probpose"):
            two_d = load_2d(config["root"], model)
            three_d = load_3d(config["root"], model)
            rows.append({"video_key": video_key, "video": config["label"], "model": model, **two_d, **three_d})
            body[video_key][model] = three_d

    plot_2d(rows, ASSET_DIR / "figure_2d_comparison.png")
    plot_knee_curves(
        {key: read_kinematic_series(config["root"], "pmpose") for key, config in RUNS.items()},
        ASSET_DIR / "figure_pmpose_knee_timeseries.png",
    )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("people_1 C3 自适应连续扩框人体数据简要统计报告")
    set_run_font(run, name="黑体", size=18, bold=True)
    add_paragraph(doc, "数据：people_1（403 对）与 people_1_near（217 对）  日期：2026-09-06", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    add_paragraph(doc, "一、摘要", size=14, bold=True, before=4, after=5)
    add_paragraph(doc, "两段视频均采用同一 C3 自适应连续扩框、最高分单人框和相同的 PMPose / ProbPose 前向。主视频中 PMPose 二维一致性和双目匹配更稳定；近距离视频中两模型二维均差都上升到约 27.5 px。", after=8)

    add_paragraph(doc, "二、二维结果", size=14, bold=True, before=4, after=5)
    table_2d = []
    for row in rows:
        table_2d.append([
            row["video"], "PMPose" if row["model"] == "pmpose" else "ProbPose",
            f"{row['mean']:.2f}", f"{row['pck25']:.2f}",
            f"{row['regions']['head']:.2f}", f"{row['regions']['shoulder_girdle']:.2f}",
            f"{row['regions']['upper_limb']:.2f}", f"{row['regions']['lower_limb']:.2f}",
        ])
    add_table(doc, ["视频", "模型", "均差\npx", "PCK25\n%", "头部\npx", "肩带\npx", "上肢\npx", "下肢\npx"], table_2d, [3.2, 2.0, 1.5, 1.6, 1.4, 1.4, 1.4, 1.4])
    add_paragraph(doc, "表 1 二维结果相对 Sapiens2 操作性参考。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(ASSET_DIR / "figure_2d_comparison.png"), width=Cm(16.5))
    add_paragraph(doc, "图 1 C3 自适应连续扩框的二维比较。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    doc.add_page_break()
    add_paragraph(doc, "三、无门限三角化和人体时序", size=14, bold=True, before=0, after=5)
    table_3d = []
    for row in rows:
        flags = row["flags"]
        high = int(flags.get("high_reprojection_error", 0))
        table_3d.append([
            row["video"], "PMPose" if row["model"] == "pmpose" else "ProbPose",
            f"{row['matched']}/{row['pairs']}", str(row["points"]),
            f"{row['lower_coverage']:.2f}", str(high),
        ])
    add_table(doc, ["视频", "模型", "匹配对", "输出 3D 点", "下肢输出\n覆盖 %", "高重投影\n质量标记"], table_3d, [3.2, 2.0, 1.8, 2.0, 2.0, 2.5])
    add_paragraph(doc, "表 2 无门限三角化结果。高重投影点保留在轨迹中并单独统计。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

    table_body = []
    metric_order = [
        ("left_thigh_length_mm", "左大腿"), ("right_thigh_length_mm", "右大腿"),
        ("left_shank_length_mm", "左小腿"), ("right_shank_length_mm", "右小腿"),
        ("left_knee_angle_deg", "左膝角"), ("right_knee_angle_deg", "右膝角"),
    ]
    for video_key, config in RUNS.items():
        for model in ("pmpose", "probpose"):
            values = body[video_key][model]["t2"]
            line = [config["label"], "PMPose" if model == "pmpose" else "ProbPose"]
            for key, _ in metric_order:
                line.append(f"{float(values[key]['median']):.1f}")
            table_body.append(line)
    add_table(doc, ["视频", "模型", "左大腿\nmm", "右大腿\nmm", "左小腿\nmm", "右小腿\nmm", "左膝角\n°", "右膝角\n°"], table_body, [2.8, 1.7, 1.5, 1.5, 1.5, 1.5, 1.4, 1.4])
    add_paragraph(doc, "表 3 帧间人体数据中位数。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(ASSET_DIR / "figure_pmpose_knee_timeseries.png"), width=Cm(13.2))
    add_paragraph(doc, "图 2 PMPose 左右膝角的原始时序。曲线未平滑，缺失帧保持断开。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    add_paragraph(doc, "四、结论", size=14, bold=True, before=4, after=5)
    add_paragraph(doc, "people_1 的 PMPose 链最稳定：二维均差 13.52 px、PCK25 为 84.14%、下肢输出覆盖 99.26%。near 视频中两模型二维均差接近，PMPose 的下肢输出覆盖仍略高于 ProbPose（91.71% 对 88.94%）。骨段中位数随视频距离变化较大，当前优先用于关节相对运动和帧间连续性分析。")

    for style_name in ("Normal",):
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
