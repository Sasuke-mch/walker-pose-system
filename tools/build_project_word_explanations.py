"""Build the two requested Chinese Word documents from verified local evidence.

Run from the repository root. Does not run fitting or modify experiment outputs.
"""
from pathlib import Path
import json
import re
import shutil

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
EXP = ROOT / 'research_records/engineering_validation/G20261002_wrist_bound_scale_v1'
OUT = ROOT / '资料/实验报告'
GUIDE = ROOT / 'docs/助步器项目完整逻辑详解_20261002.docx'
REPORT = OUT / '固定抓握约束优化与腕尺度对照实验报告_20261002.docx'


def clean(s):
    return s.replace('**', '').replace('`', '')


class Book:
    def __init__(self, title, subtitle):
        self.d = Document()
        self.tables = 0
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.54)
        sec.left_margin = sec.right_margin = Cm(2.54)
        for name in ('Normal', 'Heading 1', 'Heading 2', 'Title', 'Subtitle'):
            st = self.d.styles[name]
            st.font.name = 'Times New Roman'
            st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体' if name == 'Normal' else '黑体')
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.font.size = Pt(10.5 if name == 'Normal' else {'Heading 1': 16, 'Heading 2': 12, 'Title': 24, 'Subtitle': 12}[name])
        normal = self.d.styles['Normal'].paragraph_format
        normal.line_spacing = 1.5
        normal.space_after = Pt(5)
        normal.first_line_indent = Pt(21)
        for name in ('Heading 1', 'Heading 2'):
            f = self.d.styles[name].paragraph_format
            f.first_line_indent = Pt(0)
            f.space_before, f.space_after = Pt(14), Pt(7)
            f.keep_with_next = True
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); p._p.append(f)
        self.d.core_properties.title = title
        self.d.core_properties.subject = subtitle
        self.d.core_properties.author = 'Walker Pose System 项目记录'
        for el in self.d.styles.element.iter():
            if el.tag == qn('w:color'):
                el.set(qn('w:val'), '000000')
                for key in ('themeColor', 'themeTint', 'themeShade'):
                    el.attrib.pop(qn('w:' + key), None)
            if (el.tag in [qn('w:' + x) for x in ('top','bottom','left','right')]
                    and el.getparent().tag in (qn('w:pBdr'), qn('w:tblBorders'), qn('w:tcBorders'))):
                el.set(qn('w:color'), '000000')
                for key in ('themeColor', 'themeTint', 'themeShade'):
                    el.attrib.pop(qn('w:' + key), None)
        zoom = self.d.settings.element.find(qn('w:zoom'))
        if zoom is not None:
            zoom.set(qn('w:percent'), '100')
        self.d.add_paragraph(title, 'Title')
        self.d.add_paragraph(subtitle, 'Subtitle')
        self.p('编写日期：2026 年 10 月 2 日。代码、参数与结果以本地已保存记录为依据。本次编写文档没有重跑人体拟合。')
        self.p('阅读提示：先读直观解释，再看公式和代码。遇到“通过”，请确认它指软件测试、输入契约、单步保护还是最终质量门；这四种通过不能互相替代。')

    def h(self, text, level=1):
        self.d.add_heading(text, level)

    def p(self, text, code=False):
        p = self.d.add_paragraph(clean(text))
        if code:
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                r.font.name = 'Consolas'; r.font.size = Pt(8)
        return p

    def table(self, headers, rows, caption):
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t.autofit = True
        for c, s in zip(t.rows[0].cells, headers): c.text = str(s)
        for row in rows:
            for c, s in zip(t.add_row().cells, row): c.text = clean(str(s))
        pr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            for k, v in {'val': 'single', 'sz': '4', 'color': '000000'}.items(): el.set(qn('w:' + k), v)
            borders.append(el)
        pr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription')
        for i, row in enumerate(t.rows):
            rp = row._tr.get_or_add_trPr()
            rp.append(OxmlElement('w:cantSplit'))
            if i == 0: rp.append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                sh = OxmlElement('w:shd')
                for key, value in {'fill':'FFFFFF', 'val':'clear', 'color':'auto'}.items():
                    sh.set(qn('w:' + key), value)
                c._tc.get_or_add_tcPr().append(sh)
                for p in c.paragraphs:
                    p.paragraph_format.first_line_indent = Pt(0)
                    if i == 0 or i == len(t.rows) - 1 or len(t.rows) <= 9:
                        p.paragraph_format.keep_with_next = True
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    if i == 0 or re.fullmatch(r'[\d\s./%+−–,=-]+', p.text): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.bold = i == 0
        self.tables += 1
        p = self.p(f'表 {self.tables} {caption}')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)

    def blocks(self, text):
        for block in text.strip().split('\n\n'):
            block = block.strip()
            if block.startswith('## '): self.h(block[3:])
            elif block.startswith('### '): self.h(block[4:], 2)
            else: self.p(block)

    def source_section(self, text, heading, level=1):
        self.h(heading, level)
        lines = text.splitlines(); i = 0; in_code = False
        while i < len(lines):
            s = lines[i].strip()
            if s.startswith('```'): in_code = not in_code; i += 1; continue
            if not s: i += 1; continue
            if s.startswith('|'):
                tab = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    row = [clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch('[: -]+', x or '-') for x in row): tab.append(row)
                    i += 1
                if len(tab) > 1: self.table(tab[0], tab[1:], '实现说明对照')
                continue
            if s.startswith('###'):
                p = self.p(re.sub(r'^#+\s*[\d.]+\s*', '', s))
                p.runs[0].bold = True
                p.paragraph_format.keep_with_next = True
            elif s.startswith('#'): self.h(re.sub(r'^#+\s*', '', s), 2)
            else: self.p(lines[i].rstrip() if in_code else s, code=in_code)
            i += 1

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(path)


GUIDE_TEXT = '''
## 一、先用一个故事理解整个系统

设想有人扶着无轮助步器走路，两台鱼眼相机固定在助步器上。人向前迈步时，助步器可能不动；抬起助步器向前放时，相机也跟着动。因此，画面里一个脚踝移动了，可能是人的脚动了，也可能是相机动了，还可能是两者都动了。项目首先要把这几件事分清，再谈姿态、脚底与握持。

系统把问题分成三个层次。第一层从左右图像估计二维点，再用标定几何得到三维观测。第二层估计助步器运动阶段和相机相对地面的位姿，让不同时间的三维点能放到同一个地面坐标里。第三层用有皮肤表面的参数人体拟合观测，再研究脚底、握把与手指的几何约束。第三层目前主要是离线研究，不能把它的耗时当成实时采集程序的耗时。

你可以把三维观测理解为“相机测出来的一组带误差的钉子”，把人体模型理解为“有关节、有固定表面连接关系的橡胶人偶”。拟合就是调人偶的少数控制量，让它尽量靠近钉子。脚底、手掌和握把提供额外要求，但额外要求如果不可信，也会把人偶推到错误位置。

最重要的区别是：测量值、模型预测值、人工构造目标与显示效果是四件事。画面漂亮不等于测量准确；损失下降不等于每个身体部位都改善；程序完成也不等于实验通过。后面的每一章都围绕这几个区别展开。

## 二、系统中的数据到底是什么

### 2.1 帧、帧对、序列与窗口

一帧是一张图。一个双目帧对是同一时刻附近的左图和右图。序列是一串有顺序的帧对。窗口是从序列中截取的一段。当前腕尺度实验的 270–295 是 26 个帧对，60–100 是 41 个帧对；它们来自 448 对有序输入。程序半开区间写成 [270,296) 和 [60,101)，右端不包含。

数组的第 60 行未必是原始视频的第 60 帧。另一组 360 帧主线的真实 pair_id 是 1157–1516，开发窗索引 60–90 对应真实 pair_id 1217–1247。以前曾把视频前 448 帧与另一套 PMPose 的 448 行直接对齐，造成输入错位。因此必须沿 pair_id、图像路径和时间戳追踪，不能仅比较数组长度。

### 2.2 一个关键点的完整记录

一个二维点不仅有横纵坐标，还应有相机身份、人体身份、关节编号、置信度、是否越界、来源帧等字段。三维点还增加深度、左右重投影误差、是否通过质量门及失败原因。缺失必须明确记录，不能把没有检测到的人写成一个全 NaN 的假人体。

COCO-17 是二维观测的 17 点语义：鼻、左右眼、左右耳、左右肩、左右肘、左右腕、左右髋、左右膝、左右踝。左右始终指人的左右，不随着画面镜像或转向改编号。右膝是索引 14，不能永久删除；它应逐帧按与其他点相同的规则接受或拒绝。

## 三、采集与相机身份：先保证左右没有认错

电脑中的 OpenCV 摄像头 0、1 是当前枚举顺序，拔插设备后可能变化。项目用 PnP instance_id 注册表找到实际设备，再赋予 cam0=LEFT、cam1=RIGHT 的标定角色。如果设备身份错了，后面的标定矩阵即使数值正确，也是在解释错误的相机。

双目采集保存原图、帧编号、时间戳、读取耗时和丢帧信息。配对依据 stereo_pairs.csv 和左右帧 CSV。主机读取返回时间能用于工程配对，但不自动等于相机精确曝光时刻。人正在快速摆动时，左右不同步会把两种身体状态误当作同一时刻来求交。

360 对保存帧的既有配对记录中，左右主机时间差均值 9.35 ms、P95 16.57 ms。这说明应保留同步误差分析，不能仅凭“两张图都读到了”就宣布同步精确。本阶段腕优化没有改变采集方式，也没有重新验证硬件同步。

## 四、二维检测：为什么要转图、扩框，再转回去

### 4.1 检测人与找关节是两步

人体检测框回答“人在哪个区域”；姿态模型回答“这个区域里肩、肘、腕、髋、膝、踝在哪里”。当前主线使用 PMPose 的二维结果。模型给出的置信度是筛查和加权信息，未经独立校准不能解释成“这个点有某个百分比概率是真实位置”。换一个更大的检测器也不能自动修复标定、同步和下游目标设计问题。

摄像头安装方向可能让人横在图里。模型对正常直立人体更熟悉，所以左图可先逆时针 90°、右图先顺时针 90°再推理。这个旋转是图像预处理，不是相机物理旋转。模型输出后必须逆变换回原始鱼眼像素，标定才认识这些坐标。

举例：原图宽 W、高 H，逆时针 90° 后点坐标成为 u′=v、v′=W−1−u；恢复时 u=W−1−v′、v=u′。这里减 1 来自像素索引从 0 开始。实际代码还要恢复框、图像尺寸和全部关键点，而不只是恢复一只脚。原始分辨率是 1920×1080，旋转后的宽高会交换。

### 4.2 统一的脚部 ROI

人靠近相机时，普通人体框可能截掉脚。项目复用 foot_inclusive_box()，根据框宽占图像宽的比例连续扩展。比例小于 0.37 时保留原框；超过后用 ((r−0.37)/(1−0.37))^0.75 增长。左右、顶部、底部最大扩展比例分别为 0.20、0.08、1.30，最后按图像边界裁剪。

这不是“框越大越好”。太大的框会把背景引入模型，也会缩小人体在网络输入里的有效尺寸。连续规则避免突然跳变。当前明确要求复用现有函数，不另外建立一套看似相近的扩框公式，否则离线与实时结果会失去可比性。

## 五、双鱼眼三角化：从两条视线得到一个点

### 5.1 标定提供什么

每台相机的内参和鱼眼畸变参数描述“一个空间方向会落到哪个像素”；双目外参 R、T 描述右相机相对于左相机的位置和朝向。鱼眼边缘严重弯曲，不能直接把像素差套进普通平行针孔双目的深度公式。先用标定将原始像素转换为去畸变归一化视线，再做三角化。

左相机投影矩阵写成 [I|0]，右相机写成 [R|T]。cv2.triangulatePoints 得到齐次四维结果，再除以最后一维得到三维坐标。当前几何输出以左相机为坐标原点，长度单位 mm；进入人体拟合时转成 m。把 1000 mm 忘记除以 1000，会产生整整千倍的几何尺度错误。

理想情况下两条视线相交。真实情况下检测、标定和同步都有误差，两条线通常不能完美相交，求解得到一个折中点。随后把该点投影回两张鱼眼图，检查预测像素离原检测有多远，这个距离叫重投影误差。

### 5.2 为什么还要人物关联

左右图中的第一个人不必是同一个人。系统用共同关节的鱼眼去畸变极线代价建立左右候选关联，再按代价一对一选择。当前单人严格主线固定 max_matches=1，关联代价门为 0.05。没有可靠配对就记录失败，不复制另一侧的人凑成双目。

### 5.3 三种“有点”必须分清

严格主链按固定路线核对左右二维分数、原始边界、有限性、正深度和平均重投影误差。两个视图分数均需达到 0.25，平均重投影不超过 10 px。Stage 路径还记录 q2D=√(cL×cR)；具体消费端必须按其契约检查，不能只看到 accepted 字段就猜它采用哪套规则。

有限但高误差的 raw_xyz 可以保留给诊断，却不等于严格通过。部分历史人体拟合入口采用有限正深度的候选掩码及连续重投影权重，也不等同于严格 10 px 门。文档和显示要注明是哪一种。对于本阶段固定抓握实验，逐点保护集合来自冻结输入，不能临时改成另一套门去改善结果。

左右都没有可靠观测时，完整估计支路可以显式输出时间锚插值、外推或保持；单侧存在时可用当前视线与时间锚约束深度。这些是推断点，不是直接双目测量。它们不能反过来充当严格地面更新输入，也不能混入当前固定的 SMPL 观测。渲染时更不能再偷偷平滑一遍。

## 六、Stage 1/2：人在动，还是助步器在动

Stage 1 的工作假设是助步器和相机基本静止，人相对它运动。Stage 2 的工作假设是脚近似静止，助步器和相机移动。真实运动不一定一直符合其中一种，所以还有 transition 和 warming_up，证据不足时保留这些状态比强行二选一更诚实。

算法先排除人体区域和附着相机的结构候选，再在背景 3×4 网格中提取角点。KLT 跟踪这些点，RANSAC 剔除不一致的点，估计背景运动。背景可靠性要求至少 18 个内点、内点比例至少 0.35；用 5 帧因果窗口积累证据。左右可靠视图采用保运动融合，不能让一侧错误的零运动盖过另一侧明确运动。

Stage 1 还看补偿背景后的人体上半身运动。Stage 2 看双踝能否由背景运动解释、踝间距是否稳定。当前 Stage 1 最近 5 帧至少 3 票，Stage 2 最近 3 帧至少 2 票，部分短时证据缺失允许有限保持。投票是抗噪规则，不是人工阶段真值。

## 七、动态地面：把移动相机里的点放进固定房间

### 7.1 坐标变换的直观意义

相机坐标好比“以车为原点描述路人位置”；地面坐标好比“以房间为原点描述路人位置”。相机往前走，房间不会跟着走。因此必须知道每一帧相机在房间里的 R 和 t，按 pG=RGC·pC+tGC 变换点。人体、助步器、相机光心必须用同一帧同一个变换。

Stage 1 保持相机地面位姿，并保存最近严格双踝世界位置的中位数作锚。Stage 2 从静态背景估计旋转；本质矩阵 recoverPose 退化时用单位视线的 Wahba/Kabsch 回退。左右旋转变到同一坐标后比较，不一致超过门限则拒绝或显式保持相应状态。

### 7.2 平移怎样从脚锚推出来

若某只脚在地面中的位置 A 不动，现在相机测到它的位置 p，那么 A=Rp+t，整理得 t=A−Rp。左右脚分别算一次；两者接近时取均值，相差过大时只有明确历史预测支持才选一脚，否则拒绝。这里平移主要来自脚静止假设，不是静态背景独立测得的完整平移。

给一个纯示意数值：R=I，脚的地面位置 A=(0,0,0) m，相机观察 p=(0,0,−0.7) m，则相机平移 t=(0,0,0.7) m。这只是帮助理解符号，不是本项目某帧的测量结果。若脚实际抬高但仍被当成固定锚，算法会把脚运动的一部分错算成相机运动。

当前估计器使用因果中位数和更新增益，这属于估计方法；最终显示不再追加跨帧美化。Stage 2 结束最多允许有限 transition 求解；重新确认 Stage 1 时重锚高度、俯仰和横滚，保留累计平面运动与偏航。held、rejected、unavailable 不能统一涂成正常更新。

### 7.3 自洽循环为什么必须警惕

用脚静止推相机运动，再用推出来的相机运动证明脚静止，是同一个假设绕了一圈，不能当作独立验证。下一步需要标记板、外部刚体轨迹或其他独立位姿依据。现有内部法向漂移和切换跳变量可以发现异常，但不能给出相机轨迹真实误差。

## 八、助步器结构：它不是画出来的几根装饰线

结构文件给出节点、连接关系、扶手段和安装位姿。点先从助步器局部坐标进入相机坐标，再随当前相机地面变换进入房间。助步器安装结构与相机一起动，所以不能每帧另外手工挪动助步器去迎合人体。

目前显示约定使用实体圆柱或棱柱杆件，至少包含四脚、立柱、扶手、侧横杆和前横杆。相机与前横杆之间 30 mm 安装偏置是工程假设，不能写成现场实测。一个手腕离扶手近，只说明几何接近；是否真正握住、是否受力还需要表面几何与独立观测。

## 九、人体模型：从 17 个点到有皮肤的人

SMPL 类模型把人体表面表示成固定连接的三角网格，常用身体表面有 6890 个顶点、13776 个三角面。体型 beta 控制高矮胖瘦及身体比例，关节旋转控制姿态，根旋转和平移控制整个人在空间的位置。模型通过关节层级和蒙皮把这些参数转成全部表面顶点。

SMPL 主要提供身体；SMPL-H 增加可表达手指的手部；SMPL-X 还扩展面部等表达。仓库存在多条研究支路，不意味着它们都通过验收。当前一般身体主线保留 SMPL 与共享 beta；这轮固定抓握使用 SMPL-H、VPoser 身体和左右 MANO PCA 手参数，不应把两者写成同一个实验。

VPoser 将 32 维潜变量解码为 21 个身体关节旋转。可把潜变量想成一组组合旋钮：某个旋钮可能同时影响肩、背、腿。它不是“第一个数专门管左肘”。所以即便阶段叫 upper_only，只优化潜变量也可能动腿，必须另有腿部保护。

每只手的 PCA 用 12 个系数组合出手指关节姿态。当前窗口中同一只手共享一组系数，表达持续握持的候选形状；它不允许手指每帧随意变化。共享能减少自由度和抖动，也可能难以解释真实手形变化。这个取舍必须用观测和几何验证，不能因模型更稳定就称为更准确。

## 十、四种点与三种坐标不要混用

三角化 COCO 点是观测。模型 COCO 点由表面顶点通过 observation regressor 加权得到，用于和同语义观测比较。SMPL-H 原生关节是模型运动学节点。表面顶点描述皮肤、脚底和手掌。模型脚踝不是鞋底，模型腕不是掌心，这就是为什么接触损失要用表面。

本轮目标腕为原生关节 20/21，观测腕为 COCO 9/10。它们有数毫米偏移，不能直接说“都是腕所以一定重合”。当前短窗最大差约 5.8–6.2 mm，能解释部分语义偏差，却远不足以单独解释 70–126 mm 量级的最终腕偏差。

人体与三维观测在左相机系比较；抓握目标在助步器系比较；脚高度与脚锚在地面系比较。代码中的行向量和列向量写法看起来可能转置相反，但只要约定一致就正确。最可靠的检查是往返变换误差、旋转正交性、单位和已知点，而不是凭公式长得像不像。

## 十一、拟合到底在求什么

定义 x 为当前所有允许修改的身体和手参数，forward(x) 生成模型点、网格和旋转。每个损失项问一个问题：模型离三维观测多远，投回图像离二维点多远，腕离目标多远，姿态离先验多远，脚底是否到地面以下，手是否穿进握把。把这些惩罚乘权重加起来得到总损失 L(x)。

归一化尺度先把不同单位转成可比较的无量纲残差。例如 10 mm 腕误差除以 5 mm 得 2，平方得 4；再乘位置权重 20 得 80。改变分母与改变权重都改变相对作用。不能比较两套不同公式的 L 值大小来判断哪种姿态更好。

软损失允许交换。若腕项降低 100，腿项升高 20，总损失仍下降 80。优化器可能很认真地降低目标，同时把腿变差。硬约束则是通行证，任意一项不满足就拒绝候选。因此“加了脚损失”与“脚绝不会变差”没有等价关系。

每个接受的观测点有不可变预算：当前误差最多比最初参考多 10 mm，另保留数值容差。原误差 30 mm 的点上限约 40 mm；走到 39 mm 后不能把新上限滚成 49 mm。预算保护的是不继续显著变坏，原本 30 mm 的误差并没有因此变成零。

## 十二、脚底、静止与支撑为什么最容易被混淆

脚底接触项让候选表面靠近地面；非穿透项惩罚表面落到地面以下；切向速度项抑制明确 support 状态下沿地面滑动。这三个要求不相同。摆动脚可以在地面上方，不应被强制贴地；支撑脚应有连续的状态证据，不能只因为某一帧离地面近就判断承重。

C1/C2 路线已经修正为只在接触 Stage D 生效，避免提前改变身体初始化。切向项要同时满足相邻两帧 support 与 frame_ok；没有活动点组时，该项没有监督，不能把零值解释成没有滑动。当前保留表面=1、非穿透=1、切向=0.001 的开发配置，但数据仍缺独立 support 标签。

固定抓握实验另有 Stage 2 静止脚锚假设：从不可变完整序列的有效段选择代表脚形，对 heel/ball 的固定小片顶点建立目标，做整体 XY 对齐并抬到地面。它使用实际代表帧的脚最低点决定抬升，不是逐顶点压平，也不是每个窗口重新发明一套锚。

本轮两窗的活动脚锚帧中，冻结表面 support 候选都为 false。这说明两个模块采用不同的工程依据。不能自动宣布“Stage 2 错了”，也不能强行把 support 改成 true。应检查阶段、地面、标签来源和真实脚运动，再决定哪种假设可用。

## 十三、手与握把：腕到位还不够

把腕搬到握把附近，不保证手指包住杆，也不保证手掌没有穿杆、手指没有互相穿透。当前损失涉及握把胶囊距离、掌面接近、手指分区接近、拇指与其余手指相对方向、局部形状参考及自碰撞。它们共同描述几何候选。

握把用半径 16 mm 的胶囊近似。到胶囊表面的有符号距离为负时代表穿入。局部抓握模板独立通过包裹检查，并不保证放回完整身体后仍通过：腕部附近的蒙皮受前臂与身体关节影响，整只手未必是严格刚体。必须从最终完整人体网格重新审计。

完整手审计检查每个预期帧对的左右手，而不是只看几个截图。270 窗共 52 个手帧，60 窗共 82 个手帧。本轮两尺度分别都是 0/52 和 0/82。只要一条记录缺失、重复或来源错误，即使剩余记录通过，也不能说完整窗口通过。

## 十四、如何读指标而不被骗

平均数回答典型总体水平，P95 更关注靠后的较差部分，最大值关注最极端情况。P95=78 mm 并非每帧都是 78 mm，也不是误差绝不超过 78 mm。身体 RMS 是平方后平均再开方，大误差更有影响。脚最低 Z 是相对当前工程地面的高度，负值提示内部穿地，而不是直接测得真实鞋底陷入地面的距离。

一个总 RMS 可以被不同身体组平均掉。本轮第一窗 5 mm 的总体 RMS 改善约 0.436 mm，同时头部分组恶化约 5.641 mm，右脚最低值恶化约 4.336 mm。因此报告必须同时列身体分组、左右侧、尾部指标和几何门。

接受步数衡量搜索是否实际提交了更新。它不是准确率。初始投影认证、非线性恢复、候选 forward、最终提交是不同计数层级。一百次尝试可能只有一步被提交；多个拒绝原因还可能同时存在，不能把原因数相加当成独立步数。

## 十五、当前最好方法与没有完成的工作

如果问题是“哪种候选最值得继续开发”，已有三路线短窗对照支持 corrected：它在保留原观测门的条件下，让腕位置比 Adam 回退与单次投影更接近目标。如果问题是“哪种方法已经满足完整身体、脚和手质量”，答案是当前这些候选都没有通过。

本阶段新增软界尺度参数与诊断，并完成 1 mm/5 mm 两窗对照。5 mm 不具备全指标优势，默认保留 1 mm。A 的局部可行方向诊断和 B 的约束感知更新有实际实现；C 的尺度候选已经验证失败，但新的调度消融尚未执行；D 有来源与锚兼容性诊断，物理依据仍待补；E 时序消融和 F 新的完整身体手表面精修方案尚未作为本轮独立实验推进。

不能按字母顺序盲目堆新模块。当前日志已经显示不少候选保护可行却没有目标下降，继续增加可行恢复次数未必对症。下一步应先拆清各项梯度与真实试步变化，再预注册一个有明确预算的下降方向候选；跨帧项需要耦合优化，不能直接塞进本帧雅可比捷径。

## 十六、可视化应该帮你发现什么

合格页面应同时显示真实三角面、模型 COCO 点、原始三角化点、拒绝状态、地面、助步器实体、双相机光心、当前阶段和变换状态。骨架连接原始观测的 COCO 点，所有连线黑色实线；有限拒绝点保留叉号和原因，不能用删掉失败点的方式让画面干净。

网页采用统一显示变换 [X,−Y,Z]，Z 始终竖直。视频的指定物理视角可使用 [X,Y,Z]，应按各自规范声明。变换是显示约定，不能混进优化。三角面索引必须展平；13776 个三角面需要 41328 个索引，不能把嵌套数组直接交给 Uint32Array。

播放只切换帧，不插值网格、不冻结缺失点、不吸附脚踝。30 FPS 的 1 倍速约每 33.3 ms 前进一帧。查看一个结果时，优先对照原始双图、腕目标、脚底最低点和失败状态；不要先凭人体是否流畅判断实验成功。

## 十七、从零开始检查一个异常

若左右手突然颠倒，先查相机身份、人体左右语义、逆旋转和坐标变换，不先换网络。若三维点飞远，先查双目是否同一人同一时刻、原始边界、去畸变与正深度。若脚整段埋地，区分相机地面位姿、人体整体高度、脚底集合与接触标签问题，不直接增加脚项权重。

若 loss 不下降，先查 finite、requires_grad、梯度是否存在、参数是否真正变化、候选被哪一道门拒绝。若梯度正常而候选总越界，要研究方向；若保护已经可行但目标上升，要研究目标下降与损失取舍；若持续接受但最终手几何失败，要研究目标和验收之间的差异。

若两个实验突然一个更好，检查输入、初始化、窗口、配置、预算、脚状态、碰撞刷新、完整覆盖是否相同。不要只看文件名里的 1mm/5mm。当前独立审计正是重新读取保存参数和结果核对这些事实，而不是相信运行脚本自己写了“成功”。

## 十八、学习顺序与自测

建议先掌握像素、视线、三角化和坐标变换，再学习参数人体、损失与梯度，最后读约束优化。阅读下面实现详解时，反复问自己三个问题：变量是什么？目标是什么？哪一条失败会使这一步回滚？能回答这三个问题，就抓住了优化器主干。

自测一：只有左图一个腕点，能否唯一确定三维腕位置？不能，只有视线，深度还需其他信息。自测二：总损失下降，脚能否变差？能，只要脚不是相应硬门且别的项收益更大。自测三：5 mm 软界尺度是否把最终腕允许误差改成 5 mm？没有，最终位置门仍是 10 mm，5 mm 只是惩罚分母。

自测四：探针找到局部下降方向，是否证明目标可达？没有，它只给出某个状态的一步见证。自测五：0 个活动 support 点时切向项为 0，是否证明脚不滑？没有，因为监督没有激活。自测六：单元测试通过、手几何 0/82，是否互相矛盾？不矛盾，前者验证程序行为，后者说明当前输出质量失败。
'''


REPORT_TEXT = '''
## 一、摘要

完成腕软界尺度参数化、拒绝诊断与两窗对照。5 mm 未通过相对门，完整手几何均失败。默认保留 1 mm，继续审查目标取舍。

## 二、背景与本阶段范围

原任务是在双目观测基础上，让完整人体靠近助步器上的固定抓握目标，同时保留身体、脚和手的几何质量。原 Adam 回退只能缩小同一方向，靠近约束边界时容易停滞。前一阶段加入线性投影和非线性恢复后，腕位置改善，但腿部时序、部分身体组及手部几何仍未通过。

本阶段检查的问题是：腕越界项是否太强，以至于身体和脚被不合理交换；当前停滞究竟来自求解未认证、真实越界，还是保护可行但总目标没有下降。实验只改变软界残差分母，未同时改变目标位置、10 mm 门、观测预算、阶段步数或学习率。

本报告将前序机制修正与本轮新增实现分开。约束投影、非线性恢复和事务回滚是已建立的方法；本轮新增尺度参数、详细拒绝分类、目标/脚锚兼容性诊断、尺度基准与独立审计。编写本报告时读取既有结果，没有新增拟合。

## 三、方法与实现过程

### 3.1 输入冻结

两窗都从不可变共享手 PCA 的完整参考结果开始，不把前一次失败候选当成下一次保护预算。身体体型、标定、抓握目标、场景变换、脚状态、完整序列脚锚都冻结。实验是同源固定输入上的优化对照，没有重新采集图像或重新运行 Stage/动态地面。

窗口从完整 448 对有序输入截取。270–295 有 26 个帧对，60–100 有 41 个帧对；左右两只手均纳入全帧审计。手覆盖要求每个预期 pair 的每侧恰好一条记录，缺失、重复、额外或非法侧别全部拒绝。

### 3.2 腕软界尺度参数化

新增 --wrist-bound-scale-m，默认 0.001 m，只接受有限正数。集中函数 wrist_bound_loss 计算 r=relu(d−0.010)/s，再返回 mean(r²)+max(r²)，外层权重仍为 10。d 是原生腕关节到构造目标的欧氏距离，s 是本轮唯一变化的尺度。

当 d=100 mm 时，越界量为 90 mm。s=1 mm 时单个平方残差为 8100，s=5 mm 时为 324，恰为 1/25。函数值及对相同输入的梯度都缩小为 1/25。这个例子解释惩罚强度，不能用它推断最终腕位置必然更好。位置项仍按 5 mm 归一化，权重 20；最终腕 P95 门仍为 10 mm。

### 3.3 事务诊断修正

身体更新现在记录 before_loss、loss、loss_delta、parameters_changed、loss_finite、exact_guard_passed 与 rejection_reasons。拒绝原因允许多项并存，包括参数未动、非有限、非线性观测越界、上肢阶段腿位移越界、PCA 越界、其他精确门失败和 objective_increased。

每个候选记录加权损失分项和左右腕平均角误差；事务起点保存原加权分项，因而能计算到底哪项变好、哪项变坏。恢复求解未认证仍保留在恢复状态，不能混成真实人体越界。初始 QP 未认证也与修正 QP 未认证分别记录。

拒绝或异常时仍恢复身体参数与完整 Adam 状态。此次日志没有把失败候选强行接受，也没有把目标上升改写成目标不可达。最后手 PCA 阶段继续使用原 guarded Adam，不冒充新的完整身体耦合手精修。

### 3.4 目标与脚锚兼容性审计

新增只读 audit_grasp_target_compatibility.py。它核对目标位姿字段、旋转矩阵正交性、帧身份，输出腕位置与朝向残差、冻结脚状态、Stage 2 段、活动锚位置残差和同段相邻顶点速度。无活动锚时写 null，避免把占位索引误当成真实脚点。

建锚在完整不可变参考序列上进行。每个有效 Stage 2 段选择接近段中位脚中心的实际代表帧，固定 heel/ball 小片顶点，整体 XY 平移到参考中心，按该帧脚最低点整体抬到地面。至少 3 个有效帧才建立相应段。审计报告只选与当前窗口重叠的段，防止把全序列其他段的抬升量归到当前窗口。

这些检查能找数组身份、变换与来源问题，不能独立认证真实左右手目标、真实支撑或全局可达性。现有数据没有独立物理接触或完整三维真值，此问题留在问题分析中，不能靠修改标签消除。

### 3.5 对照与独立审计

benchmark_wrist_bound_scale.py 在运行前写入协议。两路线均为 corrected，固定恢复最多 4 次、归一化半径 0.2、每阶段 50 步、基础学习率 0.003。两窗两尺度共四次拟合，分别完成完整手审计。

audit_wrist_bound_scale.py 重新读取保存输出，核对输入、配置、帧范围、不可变观测预算、脚状态、脚锚、初始原始指标及最终原始指标。初始化比较只排除因尺度必然变化的 wrist_bound 项，不排除其他初始指标。实际配置必须匹配预注册协议，不能只比较两路线彼此相同。

相对门要求左右腕位置 P95、左右腕角 P95、身体整体与全部可用分组 RMS、双脚最低高度、腿项均不恶化，且至少一项明确改善。位置容差为 0.001 mm，角度对应 0.001°，腿项容差 1e−6；这些是数值比较容差，不是生理意义上的最小重要差异。原五道最终质量门仍单独检查。

### 3.6 中断与恢复

第一个窗口两路线完成后，第二窗口原 1 mm 运行中断且缺最终输出。续跑前确认无活动拟合进程，把原目录和日志保留为 scale1mm_interrupted，从相同原始初始化重新运行该窗两路线。没有把中间参数与新结果拼接，continuation_record.json 保留原因和范围。

### 3.7 验证范围

上一轮执行修正时，相关 60 项测试及源码编译通过，覆盖尺度值与梯度、非法参数、目标上升分类、回滚、无活动锚的 null 及帧不匹配拒绝。历史 65 项是另一套选择范围，不能和 60 相加。本次文档工作只做文件内容、数值、Word 结构与排版检查，没有重新宣称模型测试运行。

## 四、结果

四次拟合与手审计全部结束。两个窗口的输入契约和独立一致性审计通过；两窗相对不恶化门均失败，四次拟合的原始质量门均只有观测保护通过。下面按原始单位列结果，避免以不同尺度下的总损失判优。
'''


def selected(source, lo, hi):
    return source.split(f'## {lo}. ', 1)[1].split(f'## {hi}. ', 1)[0]


def main():
    source = (ROOT / 'docs/FIXED_GRASP_CORRECTED_IMPLEMENTATION_GUIDE_20261002.md').read_text(encoding='utf-8')
    protocol = json.loads((EXP / 'protocol.json').read_text(encoding='utf-8'))
    results = json.loads((EXP / 'partial_results.json').read_text(encoding='utf-8'))
    audit = json.loads((EXP / 'independent_scale_audit.json').read_text(encoding='utf-8'))
    tx = json.loads((EXP / 'transaction_diagnostic_summary.json').read_text(encoding='utf-8'))
    g = Book('助步器项目完整逻辑详解', '从双鱼眼图像到完整身体、脚底与固定抓握优化')
    g.table(['阅读位置','你将回答的问题'], [
        ['第一至四章','系统做什么；数据如何对应；相机与二维输入为何会错'],
        ['第五至八章','像素怎样变三维；相机移动怎样消除；助步器如何定位'],
        ['第九至十三章','人体怎么动；目标与约束怎么区分；脚与手怎么验收'],
        ['第十四至十八章','指标怎么读；当前方法怎么选；异常从哪里排查'],
        ['第十九章','Adam 提案、投影、非线性恢复和回滚的逐步实现'],
        ['第二十至二十一章','本阶段改了什么；结果是否通过；去哪里读代码'],
    ], '阅读路线；Word 导航窗格可按标题定位')
    g.blocks(GUIDE_TEXT)
    g.source_section('## 2. ' + selected(source, 2, 12), '十九、当前固定抓握算法逐步实现详解')
    g.p('以上实现细节中的默认软界尺度为 1 mm。参数化、完整尺度对照和当前决策见下章；不沿用旧说明中已经过期的“尚未实现”状态。')
    g.source_section('## 19. ' + source.split('## 19. ', 1)[1], '二十、本阶段修正与最新实验证据')
    g.h('二十一、代码与权威记录导航')
    rows = [
        ('采集与身份', 'realtime_app/run_stereo.py；pose_app/camera_registry.py；pose_app/stereo_camera.py', '先确认左右设备、帧对和时间戳'),
        ('二维回映', 'realtime_app/pose_app/rotation.py；tools/build_continuous_foot_inclusive_roi.py', '模型转正后回原始鱼眼坐标；统一 ROI'),
        ('严格几何', 'realtime_app/pose_app/triangulation.py；TRIANGULATION_PIPELINE_FIXED.md', '关联、正深度、边界、误差及来源'),
        ('动态地面', 'realtime_app/pose_app/realtime_stage_walker.py；dynamic_ground_pose.py；static_background_rotation.py', '状态、背景旋转、脚锚平移'),
        ('身体与手', 'realtime_app/tools/refine_body_with_constructed_grasp.py；pose_app/constrained_grasp.py；balanced_grasp.py', '前向、损失、局部约束、事务与脚锚'),
        ('尺度评估', 'realtime_app/tools/benchmark_wrist_bound_scale.py；audit_wrist_bound_scale.py；audit_grasp_target_compatibility.py', '协议、配对结果、独立审计'),
        ('规范与当前状态', 'AGENTS.md；AI_PROGRESS.md；VISUALIZATION_PIPELINE.md', '长期约束优先，实验状态以最新追加记录为准'),
    ]
    g.p('本文解释经核对的主线及当前固定抓握研究分支。仓库中存在的其他候选脚本不因文件存在就成为已验收方法。PMPose 替代检测器、全新人体模型及后续 E/F 方案需要独立实验，本文没有把它们写成已完成。')
    g.table(['模块', '仓库内路径', '阅读目的'], rows, '项目代码入口；简写文件与前项同目录')
    g.save(GUIDE)
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(GUIDE, OUT / GUIDE.name)

    r = Book('固定抓握约束优化实验报告', '腕尺度对照：实现修正、冻结协议、完整结果与下一阶段判断')
    r.blocks(REPORT_TEXT)
    labels = [('window270_295', '270–295'), ('window60_100', '60–100')]
    rows = []; angles = []; gates = []
    for key, label in labels:
        for route, scale in [('scale1mm', '1 mm'), ('scale5mm', '5 mm')]:
            a = results[key][route]; f = a['final']
            rows.append([label, scale, f"{f['wrist_error_mm']['left']['p95']:.3f} / {f['wrist_error_mm']['right']['p95']:.3f}", f"{f['body_3d_rms_mm']:.3f}", f"{f['terms']['leg_ground_acceleration']:.3f}"])
            angles.append([label, scale, f"{f['wrist_angle_deg']['left']['p95']:.3f} / {f['wrist_angle_deg']['right']['p95']:.3f}", f"{f['sole_minimum_z_mm']['left']:.3f} / {f['sole_minimum_z_mm']['right']:.3f}", f"{a['hand_passed']}/{a['hand_total']}"])
            gates.append([label, scale, '通过', '失败', '失败', '失败', '失败'])
    r.table(['窗口', '尺度', '腕位置 P95 左/右 mm', '身体 RMS mm', '腿加速度代价'], rows, '身体与腕位置；腿项是归一化代价')
    r.table(['窗口', '尺度', '腕角 P95 左/右 °', '脚最低 Z 左/右 mm', '手通过/总数'], angles, '朝向、脚高度与完整手几何')
    r.table(['窗口', '尺度', '观测', '腕', '脚', '腿', '手'], gates, '原五门结论')
    r.h('4.1 分组变化', 2)
    r.p('下表全部为 5 mm 减 1 mm。RMS 增加表示更差。不能用整体均值覆盖分组退化。')
    r.table(['身体组', '270–295 差值 mm', '60–100 差值 mm'], [[group, f"{audit['windows']['window270_295']['body_group_difference_mm'][group]:+.3f}", f"{audit['windows']['window60_100']['body_group_difference_mm'][group]:+.3f}"] for group in ['left_arm','right_arm','torso','left_leg','right_leg','head']], '身体分组观测 RMS 的配对差值')
    r.h('4.2 更新与拒绝', 2)
    rows = []
    for key, label in labels:
        for route in ('scale1mm','scale5mm'):
            a = tx[key][route]
            rows.append([label, route.replace('scale',''), f"{a['body_accepted']}/150", a['recovery_forward_evaluations'], a['all_trial_rejection_reason_counts'].get('objective_increased',0), a['last_rejected_guard_feasible']])
    r.table(['窗口','尺度','提交步','恢复 forward','目标上升尝试','末次保护可行但拒绝'], rows, '事务和候选尝试分开统计')
    r.p('四路线的初始投影未认证计数均为 0，接受步的精确保护失败计数也均为 0。恢复未认证次数分别为 130、1、40、28。目标上升尝试数不能当作独立优化步数；一笔事务会试多个步长与修正，多种拒绝原因也可能并存。')
    r.p('第一窗 5 mm 减少了观测越界候选，但目标上升拒绝更多。第二窗也有大量保护已可行的末次拒绝。这支持继续分析目标取舍与下降方向，而不支持简单把停滞归因于恢复次数太少。forward 数反映额外计算，当前没有统一计时的端到端速度对照。')
    r.h('4.3 脚锚来源与残差', 2)
    r.table(['窗口','活动帧/相邻对','相关段','建锚抬升左/右 mm'], [['270–295','17 / 16','269–286','79.820 / 60.859'],['60–100','14 / 13','68–81','18.247 / 30.295']], '与当前窗口重叠的完整序列脚锚段')
    r.table(['窗口','尺度','最大锚残差 mm','最大顶点速度 m/s'], [['270–295','1 mm','106.057','2.769'],['270–295','5 mm','99.978','2.940'],['60–100','1 mm','75.057','1.898'],['60–100','5 mm','70.444','1.898']], '活动锚的工程地面诊断')
    r.p('两个窗口活动锚帧的冻结表面 support 候选双脚均为 false。5 mm 两窗最大锚残差降低，但第一窗最大顶点速度增加。贴近一个设计目标不等于满足另一个静止指标，更不能据此证明真实脚部接触。')
    r.blocks('''
## 五、问题分析

### 5.1 尺度改变了取舍，没有整体胜出

270 窗 5 mm 的身体整体 RMS 与腿项下降，腕朝向改善，但左腕位置、头部与右脚最低高度变差。60 窗双脚最低高度改善，右腕位置略改善，左腕位置与朝向、整体身体和双腿分组变差，腿项增加。两窗腿项方向相反，不能把第一窗的好处推广到第二窗。

### 5.2 改进搜索与改善姿态不同

corrected 可以沿可行边界改变方向，并修复线性近似遗漏的真实越界。这解决了原回退的一类机制问题。当前完整损失仍允许身体、脚、腕之间交换，单步保护又主要约束观测退化。因此接受率提高和最终质量失败完全可以同时出现。

### 5.3 目标与观测本身仍有不确定性

腕目标来自构造抓握，地面来自工程估计，Stage 2 静止来自状态假设，表面 support 来自冻结候选。它们没有共享的独立物理真值。脚锚建锚有明显抬升，最终仍有较大残差；必须检查这些依据的兼容性，不能直接通过增加约束强度解决。

### 5.4 可重复性边界

新 1 mm 对照与旧 v3 初始化原始指标一致，但最终数值没有逐项重现。270 窗旧左腕 P95 为 71.451 mm，新为 74.542 mm；60 窗旧为 127.041 mm，新为 126.039 mm。没有分离确定差异来源，不应直接归因于随机性。尺度结论使用同一当前代码的新成对对照，旧结果只作历史参照。

代码表达式重用、浮点累加与活动约束切换可能影响迭代轨迹，但小型数值例子不能解释全部模型差异。本轮两窗各一次配对，没有重复运行方差，也没有外部三维准确率或物理接触评价。

### 5.5 目前不能给出的保证

线性认证检查可行残差，没有完整 KKT 最优性认证。固定步半径不是自适应信赖域。接受修正步后保留 Adam 原提案的矩，是启发式。逐帧雅可比依赖本帧局部约束，不能直接容纳跨帧时序与共享手参数约束。以上均是继续设计时必须保留的实现条件。

## 六、结论

保留尺度参数化和诊断，默认仍为 1 mm。5 mm 候选未通过两窗相对门；当前候选均未通过完整质量门，不扩展全 448 帧。

## 七、下一阶段计划与停止条件

下一阶段先分析已保存事务：对保护可行但目标上升的候选，逐项比较加权损失与腕、身体组、脚的原始变化。明确是哪个目标阻止进展后，再设计一个单因素候选。这个分析无需先重跑大窗口。

若新增身体/脚质量预算，应区分初始不可行恢复与最终质量验收，避免让已经穿地的初始化立刻无解，也不能把最终 −3 mm 门放宽求通过。若新增目标感知下降方向，继续保持不可变观测门和完整回滚。两者不要同时修改，否则无法识别收益来源。

时序消融要同时检查抖动、观测保留和动作变化，不能靠抹平动作降低加速度。完整身体下的手表面精修需要全帧双手覆盖和最终网格审计。未通过前不扩大序列，不把新候选设成默认，不用显示平滑或筛掉失败帧交付。
''')
    r.h('八、冻结配置与代码对应')
    r.table(['配置','值','说明'], [[k,str(v),'除尺度外冻结'] for k,v in protocol['config_except_scale'].items()], '协议中的完整拟合配置；实际窗口覆盖 start/stop')
    r.table(['损失项','权重','尺度或启用条件'], [
        ['body_3d','1','身体分组，50 mm 尺度'], ['body_2d','0.15','双鱼眼鲁棒重投影，100 px'],
        ['wrist_position','20','5 mm'], ['wrist_rotation','0.1','旋转矩阵弦距离，5°近零尺度'],
        ['wrist_bound','10','10 mm 越界量；分母 1/5 mm'], ['pose_anchor / vposer_prior','1 / 0.02','完整身体配置'],
        ['body_temporal / rotation_temporal','0.02 / 0.2','有效帧掩码；前者相对身体二阶差分'],
        ['leg_rotation / leg_acceleration','0.2 / 0.05','腿旋转与地面加速度'], ['upper_leg_anchor','1','只在 upper_only 生效'],
        ['root_reference','1','平移 50 mm，旋转 0.15 rad'], ['foot_contact / nonpenetration','1 / 1','10 mm 归一化；当前 balanced 零裕量'],
        ['foot_tangential','0.001','相邻 support 且 frame_ok'], ['stage2_position / velocity','1 / 0.1','10 mm / 0.05 m/s'],
        ['hand_penetration','20','3 mm；胶囊半径 16 mm'], ['hand_region / palm / opposition','3 / 2 / 10','分区、掌面与对向'],
        ['hand_self_collision','50','全帧筛查，25 步刷新'], ['hand_surface_shape / pose_reference','0.1 / 0.5','局部形状 / PCA 参考'],
    ], '当前配置的目标项；腕三项在上肢阶段共同渐进激活')
    r.p('脚非穿透损失在当前 balanced 模式使用 0 裕量，与最终脚最低 −3 mm 验收门不是同一个量。历史 C1/C2 支路曾使用 3 mm 非穿透裕量，不能混写为本轮设置。')
    r.h('九、算法实施细节与保护条件')
    for index, (lo, hi) in enumerate([(4,5),(6,7),(7,8),(8,9),(9,10),(10,11),(11,12)], 1):
        section = selected(source, lo, hi)
        title, body = section.split('\n', 1)
        r.source_section(body, f'9.{index} {title}', level=2)
    r.h('十、复现入口与数据清单')
    r.p('以下路径相对于 D:\\my_works\\walker_pose_system；协议中的输入为绝对路径。重新运行必须使用尚不存在的新输出目录，不能覆盖此次实验。')
    r.p('主源代码：realtime_app/pose_app/constrained_grasp.py；realtime_app/tools/refine_body_with_constructed_grasp.py；benchmark_wrist_bound_scale.py；audit_wrist_bound_scale.py；audit_grasp_target_compatibility.py。后四个简写文件与 refine 工具同目录。详细历史实现说明见 docs/FIXED_GRASP_CORRECTED_IMPLEMENTATION_GUIDE_20261002.md；其旧计划段以最新追加结果为准，本报告使用当前完成状态。')
    for name, path in protocol['inputs'].items():
        r.p(name + '：' + path, code=True)
    r.p('python realtime_app/tools/benchmark_wrist_bound_scale.py --reference-metadata research_records/engineering_validation/G20261001_annotated_shared_grip_v1/active_constructed_grip_v2/balanced_body_full448_v1/fit/run_metadata.json --output-root <新的不存在目录> --steps 50', code=True)
    r.table(['文件','应检查的内容'], [['protocol.json','输入、唯一变量、四阶段配置、门与停止条件'],['comparison.json / partial_results.json','完成状态、原始指标、手覆盖、门'],['independent_scale_audit.json','输入配置初始一致、实际输出与协议一致'],['transaction_diagnostic_summary.json','接受步、恢复实际前向、未认证和目标上升'],['scale1mm_target_anchor_audit.json / scale5mm_target_anchor_audit.json','逐帧目标、脚状态、活动锚与来源'],['continuation_record.json','中断保留与重新开始范围'],['各路线日志、run_metadata.json、result.npz、手部审计','命令配置、实际参数、完整几何及失败原因']], '实验输出目录 G20261002_wrist_bound_scale_v1')
    r.save(REPORT)
    for path in [GUIDE, REPORT]:
        d = Document(path)
        text = '\n'.join(p.text for p in d.paragraphs)
        print(json.dumps({'file':str(path),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'body_characters':len(text),'bytes':path.stat().st_size}, ensure_ascii=False))
    assert GUIDE.read_bytes() == (OUT / GUIDE.name).read_bytes()


if __name__ == '__main__':
    main()
